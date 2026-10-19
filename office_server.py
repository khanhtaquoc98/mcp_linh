import os
import re
from typing import Optional
try:
    from mcp.server.fastmcp import FastMCP
except ModuleNotFoundError:
    from mcp.server.mcpserver import MCPServer as FastMCP

# Khởi tạo FastMCP Server
mcp = FastMCP("Office-Automation-MCP")

# Đường dẫn tương đối tới file service_account.json nằm cùng thư mục
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SERVICE_ACCOUNT_FILE = os.path.join(BASE_DIR, "service_account.json")

SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets",
    "https://www.googleapis.com/auth/documents",
    "https://www.googleapis.com/auth/drive",
]


def _extract_google_id(url_or_id: str) -> str:
    """Tự động tách ID từ URL của Google Sheets hoặc Google Docs, hoặc trả về chính nó nếu đã là ID."""
    url_or_id = url_or_id.strip()
    match = re.search(r"/d/([a-zA-Z0-9-_]+)", url_or_id)
    if match:
        return match.group(1)
    return url_or_id


def _get_google_credentials():
    """Kiểm tra và khởi tạo thông tin xác thực Google Service Account."""
    if not os.path.exists(SERVICE_ACCOUNT_FILE):
        raise FileNotFoundError(
            f"Không tìm thấy file cấu hình xác thực Google tại '{SERVICE_ACCOUNT_FILE}'. "
            "Vui lòng tải file service_account.json từ Google Cloud Console và đặt vào cùng thư mục với office_server.py."
        )
    from google.oauth2.service_account import Credentials
    return Credentials.from_service_account_file(SERVICE_ACCOUNT_FILE, scopes=SCOPES)


# =====================================================================
# 1. WORD (.docx) TOOLS
# =====================================================================

@mcp.tool()
def read_word_file(file_path: str) -> str:
    """
    Đọc toàn bộ đoạn văn (paragraphs) và bảng biểu (tables, ngăn cách các ô bằng dấu '|') từ file Word (.docx).

    Args:
        file_path: Đường dẫn tới file .docx cần đọc.
    """
    try:
        expanded_path = os.path.abspath(os.path.expanduser(file_path))
        if not os.path.exists(expanded_path):
            return f"Lỗi: File không tồn tại tại đường dẫn '{expanded_path}'."

        from docx import Document
        doc = Document(expanded_path)
        output = []

        # Đọc các đoạn văn
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        if paragraphs:
            output.append("=== NỘI DUNG VĂN BẢN ===")
            output.extend(paragraphs)

        # Đọc các bảng biểu
        if doc.tables:
            output.append("\n=== BẢNG BIỂU ===")
            for idx, table in enumerate(doc.tables, start=1):
                output.append(f"\n--- Bảng {idx} ---")
                for row in table.rows:
                    row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    output.append(" | ".join(row_cells))

        if not output:
            return "File Word trống, không có nội dung văn bản hay bảng biểu nào."

        return "\n".join(output)

    except PermissionError:
        return (
            f"Lỗi quyền truy cập (PermissionError): Không thể đọc file '{file_path}'. "
            "Có thể file đang được mở bởi Microsoft Word hoặc bạn không có quyền đọc file này."
        )
    except Exception as e:
        return f"Lỗi không xác định khi đọc file Word '{file_path}': {str(e)}"


@mcp.tool()
def create_word_file(
    file_path: str,
    title: str,
    paragraphs: list[str],
    table_data: Optional[list[list[str]]] = None,
) -> str:
    """
    Tạo một file Word (.docx) mới với tiêu đề, danh sách các đoạn văn và bảng dữ liệu (nếu có).

    Args:
        file_path: Đường dẫn lưu file .docx mới.
        title: Tiêu đề chính của tài liệu.
        paragraphs: Danh sách các đoạn văn bản (mỗi phần tử là 1 đoạn).
        table_data: Dữ liệu bảng dạng mảng 2 chiều (dòng đầu tiên sẽ được in đậm làm tiêu đề bảng). Mặc định là None.
    """
    try:
        expanded_path = os.path.abspath(os.path.expanduser(file_path))
        parent_dir = os.path.dirname(expanded_path)
        if parent_dir and not os.path.exists(parent_dir):
            os.makedirs(parent_dir, exist_ok=True)

        from docx import Document
        doc = Document()

        if title:
            doc.add_heading(title, level=0)

        for para in paragraphs:
            doc.add_paragraph(str(para))

        if table_data and len(table_data) > 0:
            num_rows = len(table_data)
            num_cols = max(len(row) for row in table_data) if num_rows > 0 else 0
            if num_cols > 0:
                table = doc.add_table(rows=num_rows, cols=num_cols)
                table.style = "Table Grid"
                for r_idx, row in enumerate(table_data):
                    for c_idx, val in enumerate(row):
                        cell = table.cell(r_idx, c_idx)
                        cell.text = str(val)
                        if r_idx == 0 and cell.paragraphs and cell.paragraphs[0].runs:
                            cell.paragraphs[0].runs[0].bold = True

        doc.save(expanded_path)
        return f"Thành công: Đã tạo và lưu file Word tại '{expanded_path}'."

    except PermissionError:
        return (
            f"Lỗi quyền truy cập (PermissionError): Không thể ghi file '{file_path}'. "
            "Vui lòng đóng file nếu nó đang được mở trong Microsoft Word."
        )
    except Exception as e:
        return f"Lỗi khi tạo file Word '{file_path}': {str(e)}"


# =====================================================================
# 2. EXCEL (.xlsx) TOOLS
# =====================================================================

@mcp.tool()
def read_excel_sheet(
    file_path: str,
    sheet_name: Optional[str] = None,
    max_rows: int = 200,
) -> str:
    """
    Đọc dữ liệu từ file Excel (.xlsx) và trả về định dạng văn bản (ngăn cách bởi dấu '|').

    Args:
        file_path: Đường dẫn tới file .xlsx cần đọc.
        sheet_name: Tên sheet cần đọc (nếu để trống sẽ đọc sheet đầu tiên/active).
        max_rows: Số dòng tối đa cần đọc để tránh tràn bộ nhớ (mặc định 200).
    """
    try:
        expanded_path = os.path.abspath(os.path.expanduser(file_path))
        if not os.path.exists(expanded_path):
            return f"Lỗi: File Excel không tồn tại tại đường dẫn '{expanded_path}'."

        import openpyxl
        wb = openpyxl.load_workbook(expanded_path, data_only=True)

        if sheet_name:
            if sheet_name not in wb.sheetnames:
                return (
                    f"Lỗi: Không tìm thấy sheet '{sheet_name}'. "
                    f"Các sheet hiện có trong file: {', '.join(wb.sheetnames)}"
                )
            ws = wb[sheet_name]
        else:
            ws = wb.active

        output = [f"=== Sheet: {ws.title} (Tối đa {max_rows} dòng) ==="]
        row_count = 0

        for row in ws.iter_rows(values_only=True):
            if all(cell is None for cell in row):
                continue
            row_count += 1
            if row_count > max_rows:
                output.append(f"\n... (Đã cắt bớt, chỉ hiển thị {max_rows} dòng có dữ liệu đầu tiên) ...")
                break
            formatted_cells = ["" if cell is None else str(cell).strip() for cell in row]
            output.append(" | ".join(formatted_cells))

        wb.close()

        if row_count == 0:
            return f"Sheet '{ws.title}' trong file '{expanded_path}' đang trống."

        return "\n".join(output)

    except PermissionError:
        return (
            f"Lỗi quyền truy cập (PermissionError): Không thể đọc file '{file_path}'. "
            "File Excel có thể đang bị mở khóa bởi Microsoft Excel. Vui lòng đóng file và thử lại."
        )
    except Exception as e:
        return f"Lỗi khi đọc file Excel '{file_path}': {str(e)}"


@mcp.tool()
def write_excel_data(
    file_path: str,
    data: list[list[str]],
    sheet_name: str = "Sheet1",
) -> str:
    """
    Ghi mảng 2 chiều vào file Excel (.xlsx) mới, hoặc thêm nối tiếp vào cuối sheet nếu file đã tồn tại.

    Args:
        file_path: Đường dẫn tới file .xlsx.
        data: Dữ liệu mảng 2 chiều (danh sách các dòng, mỗi dòng là danh sách các giá trị ô).
        sheet_name: Tên sheet cần ghi dữ liệu (mặc định là 'Sheet1').
    """
    try:
        expanded_path = os.path.abspath(os.path.expanduser(file_path))
        parent_dir = os.path.dirname(expanded_path)
        if parent_dir and not os.path.exists(parent_dir):
            os.makedirs(parent_dir, exist_ok=True)

        import openpyxl

        if os.path.exists(expanded_path):
            wb = openpyxl.load_workbook(expanded_path)
            if sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
            else:
                ws = wb.create_sheet(title=sheet_name)
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = sheet_name

        for row in data:
            ws.append(row)

        wb.save(expanded_path)
        wb.close()
        return f"Thành công: Đã ghi {len(data)} dòng dữ liệu vào sheet '{sheet_name}' của file '{expanded_path}'."

    except PermissionError:
        return (
            f"Lỗi quyền truy cập (PermissionError): Không thể ghi vào file '{file_path}'. "
            "File đang được mở trong Microsoft Excel. Vui lòng đóng file trước khi thực hiện."
        )
    except Exception as e:
        return f"Lỗi khi ghi dữ liệu vào file Excel '{file_path}': {str(e)}"


@mcp.tool()
def fill_excel_column(
    file_path: str,
    values: list[str],
    column_letter: str = "B",
    start_row: int = 2,
    sheet_name: Optional[str] = None,
) -> str:
    """
    Điền một danh sách giá trị dọc theo một cột cố định từ dòng chỉ định.
    Giữ nguyên hoàn toàn định dạng, biểu mẫu và công thức của các ô khác trong file Excel mẫu.

    Args:
        file_path: Đường dẫn tới file Excel (.xlsx) cần điền.
        values: Danh sách các giá trị cần điền lần lượt theo chiều dọc.
        column_letter: Chữ cái tên cột cần điền (ví dụ: 'A', 'B', 'C'...). Mặc định là 'B'.
        start_row: Số thứ tự dòng bắt đầu điền (1-indexed, mặc định là dòng 2 để bỏ qua tiêu đề).
        sheet_name: Tên sheet cần điền (nếu để trống sẽ dùng sheet đang active).
    """
    try:
        expanded_path = os.path.abspath(os.path.expanduser(file_path))
        if not os.path.exists(expanded_path):
            return f"Lỗi: File Excel mẫu không tồn tại tại '{expanded_path}'."

        import openpyxl
        from openpyxl.utils import column_index_from_string

        wb = openpyxl.load_workbook(expanded_path)
        if sheet_name:
            if sheet_name not in wb.sheetnames:
                wb.close()
                return (
                    f"Lỗi: Không tìm thấy sheet '{sheet_name}'. "
                    f"Các sheet hiện có: {', '.join(wb.sheetnames)}"
                )
            ws = wb[sheet_name]
        else:
            ws = wb.active

        col_letter_clean = column_letter.strip().upper()
        try:
            col_idx = column_index_from_string(col_letter_clean)
        except ValueError:
            wb.close()
            return f"Lỗi: Tên cột '{column_letter}' không hợp lệ (ví dụ hợp lệ: 'A', 'B', 'AA')."

        for i, val in enumerate(values):
            current_row = start_row + i
            ws.cell(row=current_row, column=col_idx, value=val)

        wb.save(expanded_path)
        target_sheet = ws.title
        wb.close()

        end_row = start_row + len(values) - 1 if values else start_row
        return (
            f"Thành công: Đã điền {len(values)} giá trị vào cột {col_letter_clean} "
            f"(từ ô {col_letter_clean}{start_row} đến {col_letter_clean}{end_row}) "
            f"trên sheet '{target_sheet}' của file '{expanded_path}'."
        )

    except PermissionError:
        return (
            f"Lỗi quyền truy cập (PermissionError): File '{file_path}' đang bị khóa bởi Microsoft Excel. "
            "Vui lòng đóng file Excel trên máy tính rồi thử lại."
        )
    except Exception as e:
        return f"Lỗi khi điền cột vào file Excel '{file_path}': {str(e)}"


# =====================================================================
# 3. GOOGLE SHEETS & GOOGLE DOCS TOOLS
# =====================================================================

@mcp.tool()
def read_google_sheet(url_or_id: str, worksheet_name: Optional[str] = None) -> str:
    """
    Tự động tách ID từ URL Google Sheets và đọc toàn bộ dữ liệu thông qua gspread.

    Args:
        url_or_id: Đường dẫn URL đầy đủ hoặc Spreadsheet ID của Google Sheet.
        worksheet_name: Tên trang tính (worksheet) cụ thể cần đọc. Nếu để trống sẽ đọc trang đầu tiên.
    """
    try:
        creds = _get_google_credentials()
        import gspread

        client = gspread.authorize(creds)
        sheet_id = _extract_google_id(url_or_id)
        spreadsheet = client.open_by_key(sheet_id)

        if worksheet_name:
            ws = spreadsheet.worksheet(worksheet_name)
        else:
            ws = spreadsheet.sheet1

        rows = ws.get_all_values()
        if not rows:
            return f"Google Sheet '{spreadsheet.title}' (Sheet: '{ws.title}') hiện đang trống."

        output = [f"=== Google Sheet: {spreadsheet.title} | Trang: {ws.title} ==="]
        for row in rows:
            output.append(" | ".join(str(c) for c in row))

        return "\n".join(output)

    except FileNotFoundError as fnf_err:
        return f"Lỗi cấu hình: {str(fnf_err)}"
    except Exception as e:
        return (
            f"Lỗi khi đọc Google Sheet '{url_or_id}': {str(e)}. "
            "Hãy kiểm tra lại URL/ID và đảm bảo bạn đã chia sẻ (Share) quyền xem/chỉnh sửa file này cho email của Service Account."
        )


@mcp.tool()
def append_google_sheet(
    url_or_id: str,
    rows_data: list[list[str]],
    worksheet_name: Optional[str] = None,
) -> str:
    """
    Thêm nhiều dòng dữ liệu mới vào cuối Google Sheet với tùy chọn value_input_option='USER_ENTERED'
    (giúp Google Sheets tự nhận diện số, ngày tháng và công thức).

    Args:
        url_or_id: Đường dẫn URL đầy đủ hoặc Spreadsheet ID của Google Sheet.
        rows_data: Mảng 2 chiều chứa các dòng dữ liệu cần thêm vào cuối bảng.
        worksheet_name: Tên trang tính (worksheet) cần ghi. Nếu để trống sẽ ghi vào trang đầu tiên.
    """
    try:
        creds = _get_google_credentials()
        import gspread

        client = gspread.authorize(creds)
        sheet_id = _extract_google_id(url_or_id)
        spreadsheet = client.open_by_key(sheet_id)

        if worksheet_name:
            ws = spreadsheet.worksheet(worksheet_name)
        else:
            ws = spreadsheet.sheet1

        ws.append_rows(rows_data, value_input_option="USER_ENTERED")
        return (
            f"Thành công: Đã thêm {len(rows_data)} dòng vào cuối trang '{ws.title}' "
            f"của Google Sheet '{spreadsheet.title}'."
        )

    except FileNotFoundError as fnf_err:
        return f"Lỗi cấu hình: {str(fnf_err)}"
    except Exception as e:
        return (
            f"Lỗi khi thêm dữ liệu vào Google Sheet '{url_or_id}': {str(e)}. "
            "Hãy đảm bảo bạn đã cấp quyền Editor (Người chỉnh sửa) trên Google Sheet cho email trong service_account.json."
        )


@mcp.tool()
def read_google_doc(url_or_id: str) -> str:
    """
    Tự động tách ID từ URL Google Docs và đọc toàn bộ nội dung văn bản thông qua Google Docs API.

    Args:
        url_or_id: Đường dẫn URL đầy đủ hoặc Document ID của file Google Docs.
    """
    try:
        creds = _get_google_credentials()
        from googleapiclient.discovery import build

        doc_id = _extract_google_id(url_or_id)
        service = build("docs", "v1", credentials=creds)
        document = service.documents().get(documentId=doc_id).execute()

        title = document.get("title", "Không rõ tiêu đề")
        content = document.get("body", {}).get("content", [])

        text_chunks = [f"=== Tiêu đề Google Doc: {title} ===\n"]
        for element in content:
            if "paragraph" in element:
                elements = element["paragraph"].get("elements", [])
                for elem in elements:
                    text_run = elem.get("textRun")
                    if text_run and "content" in text_run:
                        text_chunks.append(text_run["content"])
            elif "table" in element:
                text_chunks.append("\n[Bảng dữ liệu]\n")
                for row in element["table"].get("tableRows", []):
                    row_vals = []
                    for cell in row.get("tableCells", []):
                        cell_text = []
                        for cell_content in cell.get("content", []):
                            if "paragraph" in cell_content:
                                for pe in cell_content["paragraph"].get("elements", []):
                                    tr = pe.get("textRun")
                                    if tr and "content" in tr:
                                        cell_text.append(tr["content"].strip())
                        row_vals.append(" ".join(cell_text))
                    text_chunks.append(" | ".join(row_vals) + "\n")

        return "".join(text_chunks).strip()

    except FileNotFoundError as fnf_err:
        return f"Lỗi cấu hình: {str(fnf_err)}"
    except Exception as e:
        return (
            f"Lỗi khi đọc Google Doc '{url_or_id}': {str(e)}. "
            "Hãy kiểm tra Google Docs API đã được bật trên Google Cloud Console và file đã được chia sẻ quyền xem cho Service Account chưa."
        )


@mcp.tool()
def append_google_doc(url_or_id: str, text_to_append: str) -> str:
    """
    Tự động tách ID từ URL Google Docs và chèn thêm đoạn văn bản mới vào cuối tài liệu qua Google Docs API.

    Args:
        url_or_id: Đường dẫn URL đầy đủ hoặc Document ID của file Google Docs.
        text_to_append: Nội dung văn bản cần chèn thêm vào cuối bài.
    """
    try:
        creds = _get_google_credentials()
        from googleapiclient.discovery import build

        doc_id = _extract_google_id(url_or_id)
        service = build("docs", "v1", credentials=creds)

        # Lấy độ dài hiện tại của document để xác định endIndex
        document = service.documents().get(documentId=doc_id).execute()
        title = document.get("title", "Không rõ tiêu đề")
        body_content = document.get("body", {}).get("content", [])

        end_index = 1
        if body_content:
            last_element = body_content[-1]
            end_index = last_element.get("endIndex", 1) - 1
            if end_index < 1:
                end_index = 1

        formatted_text = text_to_append if text_to_append.startswith("\n") else f"\n{text_to_append}"

        requests_body = [
            {
                "insertText": {
                    "location": {"index": end_index},
                    "text": formatted_text,
                }
            }
        ]

        service.documents().batchUpdate(
            documentId=doc_id, body={"requests": requests_body}
        ).execute()

        return f"Thành công: Đã chèn thêm văn bản vào cuối tài liệu Google Doc '{title}'."

    except FileNotFoundError as fnf_err:
        return f"Lỗi cấu hình: {str(fnf_err)}"
    except Exception as e:
        return (
            f"Lỗi khi ghi thêm vào Google Doc '{url_or_id}': {str(e)}. "
            "Hãy đảm bảo Google Docs API đã bật và tài khoản Service Account có quyền Chỉnh sửa (Editor) trên tài liệu này."
        )


if __name__ == "__main__":
    mcp.run()
